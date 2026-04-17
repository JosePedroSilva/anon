"""End-to-end: anonymise a doc → file written → mapping decryptable →
LLM-style response rehydrates back to original names.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document as _DocxDocument

spacy = pytest.importorskip("spacy")
for _model in ("en_core_web_lg", "it_core_news_lg"):
    if not spacy.util.is_package(_model):
        pytest.skip(f"spaCy model {_model} not installed", allow_module_level=True)

from docshield import cli
from docshield.core import mapping_store
from docshield.core.rehydrate import rehydrate


def _all_text(path: Path) -> str:
    d = _DocxDocument(path)
    parts = [p.text for p in d.paragraphs]
    for tbl in d.tables:
        for row in tbl.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_end_to_end_english(simple_docx: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    # Redirect audit log to tmp so we don't litter the user's data dir.
    log_target = tmp_path / "audit.log"
    monkeypatch.setattr("docshield.logging_setup.audit_log_path", lambda: log_target)
    monkeypatch.setattr("docshield.config.config_path", lambda: tmp_path / "settings.json")

    rc = cli.main(
        [
            "anonymise",
            str(simple_docx),
            "--output-dir",
            str(tmp_path),
            "--password",
            "pw",
            "--user",
            "tester",
        ]
    )
    assert rc == 0

    out_doc = tmp_path / "sample_anon.docx"
    out_map = tmp_path / "sample_mapping.dsm"
    assert out_doc.is_file()
    assert out_map.is_file()
    assert log_target.is_file()

    text = _all_text(out_doc)
    # original PII gone, placeholders present
    assert "Giulia Rossi" not in text
    assert "giulia@example.com" not in text
    assert "[PERSON_" in text
    assert "[EMAIL_" in text

    # Mapping is encrypted; correct password reveals the originals.
    mapping = mapping_store.load(out_map, password="pw")
    originals = {e["original"] for e in mapping.entries}
    assert "Giulia Rossi" in originals

    # Simulate an LLM response and rehydrate it back.
    llm_out = "[PERSON_1] should be reachable at [EMAIL_1]."
    restored = rehydrate(llm_out, mapping)
    assert "Giulia Rossi" in restored
    assert "giulia@example.com" in restored


def test_end_to_end_italian(italian_docx: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("docshield.logging_setup.audit_log_path", lambda: tmp_path / "audit.log")
    monkeypatch.setattr("docshield.config.config_path", lambda: tmp_path / "settings.json")

    rc = cli.main(
        [
            "anonymise",
            str(italian_docx),
            "--output-dir",
            str(tmp_path),
            "--password",
            "pw",
        ]
    )
    assert rc == 0
    out_doc = tmp_path / "italian_anon.docx"
    text = _all_text(out_doc)
    assert "BNCMRA80A01F205X" not in text  # codice fiscale gone
    assert "IT60X0542811101000000123456" not in text  # IBAN gone
    assert "[ID_NUMBER_" in text
    assert "[IBAN_" in text
